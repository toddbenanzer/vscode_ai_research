import sys
from pptx import Presentation

def verify(filepath):
    print(f"Verifying {filepath}...")
    try:
        prs = Presentation(filepath)
    except Exception as e:
        print(f"FAILED: Could not open PPTX: {e}")
        sys.exit(1)

    # Slide 1: Title
    slide1 = prs.slides[0]
    if slide1.shapes.title.text != "Project Refactor":
        print("FAILED: Slide 1 title mismatch")
        sys.exit(1)
    # Check subtitle
    # Placeholder 1 is usually subtitle on title slide
    if len(slide1.placeholders) > 1:
        if slide1.placeholders[1].text != "Improving Readability and Simplicity":
            print("FAILED: Slide 1 subtitle mismatch")
            sys.exit(1)

    # Slide 2: Bullets
    slide2 = prs.slides[1]
    if slide2.shapes.title.text != "Goals":
        print("FAILED: Slide 2 title mismatch")
        sys.exit(1)
    if "Refactor code" not in slide2.placeholders[1].text:
        print("FAILED: Slide 2 content mismatch")
        sys.exit(1)

    # Slide 3: 2 Columns
    slide3 = prs.slides[2]
    if slide3.shapes.title.text != "Comparison":
        print("FAILED: Slide 3 title mismatch")
        sys.exit(1)
    # Check for text boxes (not placeholders)
    text_found = False
    for shape in slide3.shapes:
        if shape.has_text_frame:
            if "Complex logic" in shape.text_frame.text:
                text_found = True
    if not text_found:
        print("FAILED: Slide 3 column content not found")
        sys.exit(1)

    # Slide 4: Visual Request
    slide4 = prs.slides[3]
    visual_found = False
    for shape in slide4.shapes:
        if shape.has_text_frame and "[VISUAL DATA REQUEST]" in shape.text_frame.text:
            visual_found = True
    if not visual_found:
        print("FAILED: Slide 4 visual request not found")
        sys.exit(1)

    print("SUCCESS: Verification passed.")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python verify_pptx.py <pptx_file>")
        sys.exit(1)
    verify(sys.argv[1])
