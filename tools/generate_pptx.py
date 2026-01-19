import json
import argparse
import logging
import sys
from pathlib import Path
from typing import List, Callable, Dict, Any

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Configure logging
logger = logging.getLogger(__name__)

# --- CONSTANTS: JSON Keys ---
KEY_LAYOUT = "layout"
KEY_TITLE = "title"
KEY_SUBTITLE = "subtitle"
KEY_CONTENT = "content"
KEY_VISUAL = "visual_data_description"
KEY_NOTES = "speaker_notes"
KEY_CLASSIFICATION = "classification"
KEY_CONFIDENCE = "data_confidence"
KEY_SOURCES = "source_references"
KEY_HEADER = "header"
KEY_BULLETS = "bullets"

# --- CONSTANTS: Layout Names ---
LAYOUT_TITLE = "title"
LAYOUT_SECTION = "section"
LAYOUT_BULLETED = "bulleted"
LAYOUT_2_COL = "2_col"
LAYOUT_3_COL = "3_col"
LAYOUT_4_COL = "4_col"

# --- CONSTANTS: Defaults ---
DEFAULT_LAYOUT = LAYOUT_BULLETED
DEFAULT_CLASSIFICATION = "Internal Use Only"
DEFAULT_CONFIDENCE = "Unknown"
COLUMN_LAYOUTS = {LAYOUT_2_COL, LAYOUT_3_COL, LAYOUT_4_COL}

# --- TYPES ---
SlideInput = Dict[str, Any]
SlideContent = Dict[str, Any]

# Standard Layout Indices (Default Template)
LAYOUT_MAP = {
    LAYOUT_TITLE: 0,      # Title Slide
    LAYOUT_SECTION: 2,    # Section Header
    LAYOUT_BULLETED: 1,   # Title and Content
    LAYOUT_2_COL: 5,      # Title Only
    LAYOUT_3_COL: 5,
    LAYOUT_4_COL: 5
}

# Dimensions & Fonts
MARGIN_HORIZONTAL = Inches(0.5)
MARGIN_TOP = Inches(2.0)  # Room for Title
MARGIN_BOTTOM = Inches(1.0)
GAP_WIDTH = Inches(0.2)

VISUAL_WIDTH = Inches(4)
VISUAL_HEIGHT = Inches(1.5)
VISUAL_OFFSET = Inches(0.5)

FONT_SIZE_HEADER = Pt(18)
FONT_SIZE_BODY = Pt(14)
FONT_SIZE_VISUAL = Pt(12)
BULLET_CHAR = "• "

# Visual Placeholder Colors
COLOR_GRAY_FILL = RGBColor(200, 200, 200)
COLOR_GRAY_LINE = RGBColor(100, 100, 100)


def build_presentation(json_data: List[SlideInput]) -> Presentation:
    """
    Builds a PowerPoint presentation object from the provided JSON data.
    Does not save the file to disk.
    """
    prs = Presentation()

    for slide_data in json_data:
        layout_name = slide_data.get(KEY_LAYOUT, DEFAULT_LAYOUT)
        # Default to standard layout index if unknown
        layout_idx = LAYOUT_MAP.get(layout_name, LAYOUT_MAP[DEFAULT_LAYOUT])

        slide_layout = prs.slide_layouts[layout_idx]
        slide = prs.slides.add_slide(slide_layout)

        # 1. Common: Title
        if slide.shapes.title:
            slide.shapes.title.text = slide_data.get(KEY_TITLE, "")

        # 2. Dispatch: Content
        handler = _get_layout_handler(layout_name)
        handler(slide, slide_data, prs)

        # 3. Common: Visual Data
        visual_desc = slide_data.get(KEY_VISUAL)
        if visual_desc:
            _add_visual_placeholder(slide, visual_desc, prs.slide_width, prs.slide_height)

        # 4. Common: Notes & Metadata
        _add_notes_and_metadata(slide, slide_data)

    return prs


def _get_layout_handler(layout_name: str) -> Callable:
    """Returns the appropriate content handler for the given layout."""
    if layout_name == LAYOUT_TITLE:
        return _handle_title_subtitle
    elif layout_name == LAYOUT_BULLETED:
        return _handle_standard_bullets
    elif layout_name in COLUMN_LAYOUTS:
        return _handle_columns
    else:
        return _handle_noop


def _handle_title_subtitle(slide, data: SlideInput, prs: Presentation):
    """Handles subtitle for title slides."""
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = data.get(KEY_SUBTITLE, "")


def _handle_standard_bullets(slide, data: SlideInput, prs: Presentation):
    """Handles standard bulleted content."""
    content = data.get(KEY_CONTENT, [])
    if not content or len(slide.placeholders) < 2:
        return

    tf = slide.placeholders[1].text_frame
    tf.clear()

    item = content[0]
    for bullet in item.get(KEY_BULLETS, []):
        _add_paragraph(tf, bullet, level=0)


def _handle_columns(slide, data: SlideInput, prs: Presentation):
    """Handles multi-column content."""
    layout_name = data.get(KEY_LAYOUT, "")
    try:
        num_cols = int(layout_name.split("_")[0])
    except (ValueError, IndexError):
        num_cols = 2 # Fallback

    content = data.get(KEY_CONTENT, [])
    _add_columns(slide, content, num_cols, prs.slide_width, prs.slide_height)


def _handle_noop(slide, data: SlideInput, prs: Presentation):
    """Do nothing content handler."""
    pass


def _add_columns(slide, content: List[SlideContent], num_cols: int, slide_width: int, slide_height: int):
    """Programmatically creates text boxes for multi-column layouts."""
    col_width = _calculate_column_width(slide_width, num_cols)

    for i in range(min(num_cols, len(content))):
        left = MARGIN_HORIZONTAL + (i * (col_width + GAP_WIDTH))
        _create_column_content(slide, left, col_width, slide_height, content[i])


def _calculate_column_width(slide_width, num_cols):
    available_width = slide_width - (2 * MARGIN_HORIZONTAL)
    total_gap_width = GAP_WIDTH * (num_cols - 1)
    return (available_width - total_gap_width) / num_cols


def _create_column_content(slide, left, width, slide_height, col_data):
    top = MARGIN_TOP
    height = slide_height - MARGIN_TOP - MARGIN_BOTTOM

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    # Header (Bold)
    header_text = col_data.get(KEY_HEADER, "")
    if header_text:
        _add_paragraph(tf, header_text, bold=True, size=FONT_SIZE_HEADER, space_after=Pt(10))

    # Bullets
    for bullet in col_data.get(KEY_BULLETS, []):
        _add_paragraph(tf, f"{BULLET_CHAR}{bullet}", size=FONT_SIZE_BODY)


def _add_paragraph(text_frame, text: str, level: int = 0, bold: bool = False, size: Pt = None, space_after: Pt = None):
    """
    Adds a paragraph to the text frame.
    Uses the first paragraph if it exists and is empty to avoid leading whitespace.
    """
    if len(text_frame.paragraphs) == 1 and not text_frame.paragraphs[0].text:
        p = text_frame.paragraphs[0]
    else:
        p = text_frame.add_paragraph()

    p.text = text
    p.level = level

    if bold:
        p.font.bold = True
    if size:
        p.font.size = size
    if space_after:
        p.space_after = space_after


def _add_visual_placeholder(slide, description: str, slide_width: int, slide_height: int):
    """Adds a placeholder shape for visual data requests."""
    left = slide_width - VISUAL_WIDTH - VISUAL_OFFSET
    top = slide_height - VISUAL_HEIGHT - VISUAL_OFFSET

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, VISUAL_WIDTH, VISUAL_HEIGHT)

    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_GRAY_FILL
    shape.line.color.rgb = COLOR_GRAY_LINE

    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.text = f"[VISUAL DATA REQUEST]\n{description}"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = FONT_SIZE_VISUAL
    p.font.bold = True


def _add_notes_and_metadata(slide, data: SlideInput):
    """Appends speaker notes and governance metadata to the Notes slide."""
    text_frame = slide.notes_slide.notes_text_frame

    parts = []

    # Preserve existing notes if any
    if text_frame.text.strip():
        parts.append(text_frame.text)

    # User Notes
    notes = data.get(KEY_NOTES, "")
    if notes:
        parts.append(notes)

    # Governance Metadata
    classification = data.get(KEY_CLASSIFICATION, DEFAULT_CLASSIFICATION)
    confidence = data.get(KEY_CONFIDENCE, DEFAULT_CONFIDENCE)
    sources = ", ".join(data.get(KEY_SOURCES, []))

    metadata = (
        f"--- GOVERNANCE METADATA ---\n"
        f"Classification: {classification}\n"
        f"Confidence: {confidence}\n"
        f"Sources: {sources}"
    )
    parts.append(metadata)

    text_frame.text = "\n\n".join(parts)


def _parse_args():
    parser = argparse.ArgumentParser(description="Generate PowerPoint from JSON source.")
    parser.add_argument("input_json", type=Path, help="Path to input JSON file.")
    parser.add_argument("--output", "-o", type=Path, help="Output path (default: input_filename.pptx)")
    return parser.parse_args()


def _load_json(path: Path) -> List[SlideInput]:
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def main():
    # Configure logging only when running as a script
    logging.basicConfig(level=logging.INFO, format='%(message)s')

    args = _parse_args()

    if not args.input_json.exists():
        logger.error(f"File not found: {args.input_json}")
        sys.exit(1)

    output_path = args.output if args.output else args.input_json.with_suffix(".pptx")

    try:
        data = _load_json(args.input_json)
        prs = build_presentation(data)
        prs.save(output_path)
        logger.info(f"Presentation saved to: {output_path}")

    except json.JSONDecodeError as e:
        logger.error(f"Invalid JSON format: {e}")
        sys.exit(1)
    except Exception as e:
        logger.error(f"An error occurred: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()
