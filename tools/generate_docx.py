import json
import argparse
import logging
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor

# Configure logging
logger = logging.getLogger(__name__)

# --- TYPES ---
SlideData = dict[str, Any]

# --- CONSTANTS ---
# JSON Keys
KEY_TITLE = "title"
KEY_SUBTITLE = "subtitle"
KEY_LAYOUT = "layout"
KEY_SPEAKER_NOTES = "speaker_notes"
KEY_CONTENT = "content"
KEY_HEADER = "header"
KEY_BULLETS = "bullets"
KEY_VISUAL_DESC = "visual_data_description"
KEY_CLASSIFICATION = "classification"
KEY_CONFIDENCE = "data_confidence"
KEY_SOURCES = "source_references"

# Styles
STYLE_TITLE = 'Heading 1'
STYLE_SUBTITLE = 'Subtitle'
STYLE_HEADING_2 = 'Heading 2'
STYLE_BULLET = 'List Bullet'
STYLE_NORMAL = 'Normal'

# Defaults & Formatting
DEFAULT_CLASSIFICATION = "Internal Use Only"
DEFAULT_CONFIDENCE = "Unknown"
DIVIDER_LINE = "_" * 20

TEXT_KEY_POINTS = "Key Points"
TEXT_VISUAL_PREFIX = "[VISUAL DATA REQUEST: {}]"
COLOR_GRAY = RGBColor(100, 100, 100)


def build_document(json_data: list[SlideData]) -> Document:
    """
    Builds a Word document object from the provided JSON data.
    Follows 'Report Style': Speaker Notes are body text, Slides are highlights.
    Does not save the file to disk.
    """
    doc = Document()

    for slide_index, slide_data in enumerate(json_data):
        _add_slide_section(doc, slide_data, slide_index)

        # Add page break between sections
        if slide_index < len(json_data) - 1:
            doc.add_page_break()

    return doc


def _add_slide_section(doc: Document, slide_data: SlideData, index: int) -> None:
    """Orchestrates adding a single slide's content to the document."""
    _add_title(doc, slide_data, index)
    _add_speaker_notes(doc, slide_data)
    _add_key_points(doc, slide_data)
    _add_visual_placeholder(doc, slide_data)
    _add_metadata(doc, slide_data)


def _add_title(doc: Document, slide_data: SlideData, index: int) -> None:
    """Adds the slide title and optional subtitle."""
    title = slide_data.get(KEY_TITLE, f"Slide {index + 1}")
    layout = slide_data.get(KEY_LAYOUT, "")

    doc.add_heading(title, level=1)

    if layout == "title":
        subtitle = slide_data.get(KEY_SUBTITLE)
        if subtitle:
            doc.add_paragraph(subtitle, style=STYLE_SUBTITLE)


def _add_speaker_notes(doc: Document, slide_data: SlideData) -> None:
    """Adds speaker notes as the main body text."""
    notes = slide_data.get(KEY_SPEAKER_NOTES, "")
    if notes:
        for paragraph in notes.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())


def _add_key_points(doc: Document, slide_data: SlideData) -> None:
    """Adds slide content (bullets/columns) as a highlighted list section."""
    content = slide_data.get(KEY_CONTENT, [])
    if not content:
        return

    doc.add_heading(TEXT_KEY_POINTS, level=2)

    for item in content:
        header = item.get(KEY_HEADER)
        if header:
            _add_styled_paragraph(doc, header, bold=True)

        for bullet in item.get(KEY_BULLETS, []):
            doc.add_paragraph(bullet, style=STYLE_BULLET)


def _add_visual_placeholder(doc: Document, slide_data: SlideData) -> None:
    """Adds a styled placeholder for requested visual data."""
    visual_desc = slide_data.get(KEY_VISUAL_DESC)
    if visual_desc:
        text = TEXT_VISUAL_PREFIX.format(visual_desc)
        _add_styled_paragraph(doc, text, italic=True, color=COLOR_GRAY)


def _add_metadata(doc: Document, slide_data: SlideData) -> None:
    """Adds governance metadata at the bottom of the section."""
    classification = slide_data.get(KEY_CLASSIFICATION, DEFAULT_CLASSIFICATION)
    confidence = slide_data.get(KEY_CONFIDENCE, DEFAULT_CONFIDENCE)
    sources = slide_data.get(KEY_SOURCES, [])
    source_text = ", ".join(sources) if sources else "None"

    doc.add_paragraph(DIVIDER_LINE)

    meta_text = f"Classification: {classification} | Confidence: {confidence} | Sources: {source_text}"
    _add_styled_paragraph(doc, meta_text, size=Pt(9))


def _add_styled_paragraph(doc: Document, text: str, style: str = None,
                          bold: bool = False, italic: bool = False,
                          color: RGBColor = None, size: Pt = None) -> None:
    """Helper to add a paragraph with specific formatting."""
    p = doc.add_paragraph(style=style)
    runner = p.add_run(text)

    if bold:
        runner.bold = True
    if italic:
        runner.italic = True
    if color:
        runner.font.color.rgb = color
    if size:
        runner.font.size = size


def main() -> None:
    # Configure logging here to avoid import side effects
    logging.basicConfig(level=logging.INFO, format='%(message)s')

    parser = argparse.ArgumentParser(description="Generate Word Document from JSON source.")
    parser.add_argument("input_json", type=Path, help="Path to input JSON file.")
    parser.add_argument("--output", "-o", type=Path, help="Output path (default: input_filename.docx)")

    args = parser.parse_args()

    if not args.input_json.exists():
        logger.error(f"File not found: {args.input_json}")
        exit(1)

    # Determine output path
    output_path = args.output if args.output else args.input_json.with_suffix(".docx")

    try:
        with open(args.input_json, 'r', encoding='utf-8') as f:
            data = json.load(f)

        doc = build_document(data)
        doc.save(output_path)
        logger.info(f"Document saved to: {output_path}")

    except json.JSONDecodeError as e:
        logger.error(f"Invalid JSON format: {e}")
        exit(1)
    except Exception as e:
        logger.error(f"An error occurred: {e}")
        exit(1)

if __name__ == "__main__":
    main()
