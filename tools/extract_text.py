import argparse
import logging
from pathlib import Path

from pptx import Presentation
from docx import Document
from pypdf import PdfReader
from bs4 import BeautifulSoup

__all__ = ['extract_text']

logger = logging.getLogger(__name__)

def _extract_from_pptx(filepath: Path) -> str:
    """Extracts text from slides and speaker notes in a PowerPoint file."""
    text_content = []
    prs = Presentation(filepath)
    for i, slide in enumerate(prs.slides):
        slide_text = []
        slide_text.append(f"--- Slide {i+1} ---")

        # Extract text from shapes
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_text.append(shape.text_frame.text)

        # Extract text from notes
        if slide.has_notes_slide:
            text_frame = slide.notes_slide.notes_text_frame
            if text_frame and text_frame.text.strip():
                slide_text.append("\n[Speaker Notes]:")
                slide_text.append(text_frame.text)

        text_content.append("\n".join(slide_text))

    return "\n\n".join(text_content)

def _extract_from_docx(filepath: Path) -> str:
    """Extracts text from a Word document."""
    doc = Document(filepath)
    return "\n".join(para.text for para in doc.paragraphs if para.text.strip())

def _extract_from_pdf(filepath: Path) -> str:
    """Extracts text from a PDF file."""
    text_content = []
    reader = PdfReader(filepath)
    for i, page in enumerate(reader.pages):
        page_text = []
        page_text.append(f"--- Page {i+1} ---")

        text = page.extract_text()
        if text and text.strip():
            page_text.append(text)

        text_content.append("\n".join(page_text))

    return "\n\n".join(text_content)

def _extract_from_html(filepath: Path) -> str:
    """Extracts text from an HTML file."""
    with filepath.open('r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')
        # remove scripts and styles
        for script in soup(["script", "style"]):
            script.decompose()
        return soup.get_text(separator='\n', strip=True)

# Map extensions to extractor functions
_SUPPORTED_EXTENSIONS = {
    '.pptx': _extract_from_pptx,
    '.docx': _extract_from_docx,
    '.pdf': _extract_from_pdf,
    '.html': _extract_from_html
}

def extract_text(filepath: Path) -> str:
    """
    Extracts text from the given file based on its extension.
    Raises ValueError if the file extension is not supported.
    """
    ext = filepath.suffix.lower()
    if ext not in _SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    return _SUPPORTED_EXTENSIONS[ext](filepath)

def process_file(filepath: Path):
    """Process a single file."""
    if filepath.suffix.lower() not in _SUPPORTED_EXTENSIONS:
        return

    logger.info(f"Processing: {filepath}")
    try:
        extracted_text = extract_text(filepath)
        output_path = filepath.with_suffix('.txt')
        output_path.write_text(extracted_text, encoding='utf-8')
        logger.info(f"Saved to: {output_path}")
    except Exception as e:
        logger.error(f"Error reading {filepath}: {e}")

def process_directory(directory: Path):
    """Scans the directory for supported files and extracts text."""
    logger.info(f"Scanning {directory}...")

    count = 0
    # Recursive scan
    for filepath in directory.rglob('*'):
        # Skip hidden files and directories
        if any(part.startswith('.') for part in filepath.parts):
            continue

        if filepath.is_file() and filepath.suffix.lower() in _SUPPORTED_EXTENSIONS:
            process_file(filepath)
            count += 1

    logger.info(f"Finished processing. Extracted text from {count} files.")

def main():
    # Configure logging only when running as a script
    logging.basicConfig(level=logging.INFO, format='%(message)s')

    parser = argparse.ArgumentParser(description="Extract raw text from documents (PPTX, DOCX, PDF, HTML).")
    parser.add_argument("path", nargs="?", default=Path("workspace/sources"), type=Path, help="File or Directory to scan (default: workspace/sources)")
    args = parser.parse_args()

    if not args.path.exists():
        logger.error(f"Path not found: {args.path}")
        exit(1)

    if args.path.is_file():
        process_file(args.path)
    elif args.path.is_dir():
        process_directory(args.path)
    else:
        logger.error(f"Invalid path type: {args.path}")
        exit(1)

if __name__ == "__main__":
    main()
